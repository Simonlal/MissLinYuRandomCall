from PyQt5.Qt import *
import random
import openpyxl
import docx


class Rand(QThread):
    # 结束信号，将随机的内容传递
    finishSignal = pyqtSignal(str, str)

    def __init__(self, parent=None):
        super(Rand, self).__init__(parent=parent)
        self.NameListPath = None
        self.ContentPath = None

    def DataPrepare(self, NameListPath, ContentPath):
        self.NameListPath = NameListPath
        self.ContentPath = ContentPath
        # print(self.NameListPath,self.ContentPath)

    def run(self):
        # 随机种子
        random.seed()
        # 取名字
        WorkBook = openpyxl.load_workbook(filename=self.NameListPath)
        # 获取工作簿对象
        sheet = WorkBook["name"]

        Names = []
        for cell in sheet["A"]:
            Names.append(cell.value)
        Names = Names[1:]
        PersonNum = len(Names)
        NameIndex = random.randint(0, PersonNum-1)
        RandomName = Names[NameIndex]

        # 取题目
        doc = docx.Document(docx=self.ContentPath)
        Texts = []
        for para in doc.paragraphs:
            txt = para.text
            if txt != "":
                Texts.append(txt)
        Texts = Texts[1:]
        ContentNum = len(Texts)
        ContentIndex = random.randint(0, ContentNum-1)
        RandomContent = Texts[ContentIndex]

        self.finishSignal.emit(RandomName, RandomContent)
